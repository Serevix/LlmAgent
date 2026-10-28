"""验证生成工具的模型选择、可信参数注入、文件权限、格式及外部接口约定。"""

import asyncio
from collections import Counter
from io import BytesIO
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

from docx import Document
from httpx import ASGITransport, AsyncClient
import httpx
from langchain_core.messages import AIMessage, AIMessageChunk, HumanMessage
from langgraph.checkpoint.memory import InMemorySaver
from PIL import Image
import pytest
from sqlalchemy.ext.asyncio import create_async_engine
from sqlmodel import SQLModel, select
from sqlmodel.ext.asyncio.session import AsyncSession

from app.agents.artifact_planner import may_need_artifact, plan_artifacts
from app.agents.artifact_tools import ARTIFACT_TOOLS, generate_word
from app.agents.chat_graph import ChatContext, build_chat_graph
from app.agents.context import ContextLimits
from app.artifacts.errors import ArtifactError, ArtifactProviderError
from app.artifacts.images import generate_image_bytes, validate_image_url
from app.artifacts.word import word_bytes
from app.core.config import settings
from app.infrastructure.database.database import get_session
from app.main import app
from app.models.generated_artifact import GeneratedArtifact
from app.models.conversation import Conversation
from app.rag.grounding import GroundingDecision
from app.rag.types import BookChunk, RetrievedChunk
from app.services.artifact_service import ArtifactService
from app.services.knowledge_service import KnowledgeService


SOURCE = {"book_id": "b" * 32, "title": "测试教材", "page": 4, "chunk_id": "c" * 32,
          "score": .9, "excerpt": "Checkpoint 保存图状态。"}


def png_bytes():
    """创建极小的合法图片，测试中不调用收费图片接口。"""
    output = BytesIO()
    Image.new('RGB', (24, 24), 'blue').save(output, format='PNG')
    return output.getvalue()


def test_word_preserves_markdown_structure_and_book_references():
    """检查真实 docx 中的标题、强调、代码、列表、表格及参考位置。"""
    content = '# 原理\n\n**重点**与`code` [1]。\n\n- 甲\n- 乙\n\n| 功能 | 作用 |\n|---|---|\n| Checkpoint | 保存状态 |\n\n```python\nprint(1)\n```\n\n控制字符\x01应清理'
    doc = Document(BytesIO(word_bytes('学习笔记', content, [SOURCE])))
    assert doc.paragraphs[0].text == '学习笔记'
    assert doc.paragraphs[0].style.name == 'Title'
    assert any(p.style.name == 'Heading 1' and p.text == '原理' for p in doc.paragraphs)
    assert any(run.bold and run.text == '重点' for p in doc.paragraphs for run in p.runs)
    assert any(p.style.name == 'List Bullet' and p.text == '甲' for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 1).text == '保存状态'
    assert any(p.text == 'print(1)' for p in doc.paragraphs)
    assert any('测试教材 · 第 4 页' in p.text for p in doc.paragraphs)
    assert '\x01' not in '\n'.join(p.text for p in doc.paragraphs)


def test_tool_schema_hides_identity_paths_and_verified_body():
    """模型只填写业务参数，无法指定用户、会话、幂等键或替换 Word 正文。"""
    assert set(ARTIFACT_TOOLS['generate_image'].tool_call_schema.model_fields) == {'prompt', 'title', 'aspect'}
    assert set(ARTIFACT_TOOLS['generate_word'].tool_call_schema.model_fields) == {'title', 'content_source', 'question'}


@pytest.mark.parametrize('question,expected', [
    ('Checkpoint 是什么？', False), ('请生成一张图片，画一只猫。', True),
    ('把刚才的回答导出为 Word。', True), ('请根据知识库生成 RAG 学习文档。', True),
    ('帮我画一只小猫', True),
])
def test_only_creation_candidates_enter_tool_planner(question, expected):
    """普通问题不会额外调用规划模型；生成候选仍由模型决定是否使用工具。"""
    assert may_need_artifact(question) is expected


class ToolModel:
    """模拟原生 tool_calls 和独立的依据核对，记录每种任务的执行次数。"""

    def __init__(self, calls, *, supported=True):
        """配置模型要选择的工具，以及书籍依据核对是否通过。"""
        self.tool_calls = calls
        self.counts = Counter()
        self.supported = supported

    def with_structured_output(self, schema, **kwargs):
        """返回单独的核对包装，保持模型选工具与核对结果各自独立。"""
        owner = self

        class Review:
            async def ainvoke(self, prompt, **kwargs):
                """记录一次核对请求并返回受控判断。"""
                owner.counts['review'] += 1
                return GroundingDecision(supported=owner.supported)
        return Review()

    def bind_tools(self, tools, **kwargs):
        """确认注册了真实的两个工具，并采用模型自主选择模式。"""
        assert {tool.name for tool in tools} == set(ARTIFACT_TOOLS)
        assert kwargs['tool_choice'] == 'auto'
        return self

    async def ainvoke(self, messages, **kwargs):
        """一次调用返回可选工具列表，不负责执行工具。"""
        self.counts['planner'] += 1
        return AIMessage(content='', tool_calls=self.tool_calls)

    async def astream(self, messages, **kwargs):
        """生成正常的带引用草稿，供图中的依据核对处理。"""
        self.counts['answer'] += 1
        yield AIMessageChunk(content='Checkpoint 保存图状态 [1]。')


def call(name, **args):
    """构造 LangChain 接受的工具调用结构。"""
    return {'name': name, 'args': args, 'id': 'test-call', 'type': 'tool_call'}


def test_planner_rejects_model_supplied_identity():
    """即使模型返回隐藏参数，也必须在执行前拒绝。"""
    model = ToolModel([call('generate_image', prompt='蓝色的小猫', title='猫', user_id=999)])
    with pytest.raises(ArtifactError, match='不支持的工具参数'):
        asyncio.run(plan_artifacts(model, [HumanMessage(content='请生成一张图片')], '', ContextLimits.from_settings()))


def test_image_request_uses_model_tool_choice_without_book_retrieval(monkeypatch):
    """创作图片可以独立执行，认证身份由后端注入，结果保存在消息附件中。"""
    received = []

    async def create(self, **kwargs):
        """记录工具真实传入的身份和参数，返回公开附件。"""
        received.append(kwargs)
        return {'id': 'f' * 32, 'kind': kwargs['kind'], 'title': kwargs['title'], 'filename': '猫.png',
                'media_type': 'image/png', 'size_bytes': 100, 'download_url': '/api/v1/files/test/content'}

    monkeypatch.setattr(ArtifactService, 'create', create)
    search = AsyncMock(side_effect=AssertionError('图片创作不应进入书籍检索'))
    monkeypatch.setattr(KnowledgeService, 'search', search)
    model = ToolModel([call('generate_image', prompt='蓝色背景的小猫插画', title='小猫')])

    async def run():
        """消费真实图事件并读取 checkpoint，确认一问一答及附件持久化。"""
        graph = build_chat_graph(model, InMemorySaver(), None)
        config = {'configurable': {'thread_id': 'owned-thread'}}
        result = await graph.ainvoke({'messages': [HumanMessage(content='请生成一张小猫图片')]},
                                    config=config, context=ChatContext(user_id=7, request_id='trusted-request'))
        assert len(result['messages']) == 2
        assert result['messages'][-1].additional_kwargs['artifacts'][0]['kind'] == 'image'
        assert result['artifact_calls'] == []
        assert received[0]['user_id'] == 7 and received[0]['thread_id'] == 'owned-thread'
        assert received[0]['request_key'] == 'trusted-request'
        assert model.counts == {'planner': 1} and search.await_count == 0
        restored = await graph.aget_state(config)
        assert restored.values['messages'][-1].additional_kwargs['artifacts']
        # 下一轮普通问答不重用上一轮的生成计划，也不重复生成图片。
        monkeypatch.setattr(KnowledgeService, 'search', AsyncMock(return_value=[]))
        again = await graph.ainvoke({'messages': [HumanMessage(content='什么是 Milvus？')]},
                                   config=config, context=ChatContext(user_id=7))
        assert len(received) == 1 and len(again['messages']) == 4
    asyncio.run(run())


@pytest.mark.parametrize('supported', [True, False])
def test_new_word_content_must_pass_book_grounding(monkeypatch, supported):
    """新主题文档先走真实检索和核对，失败时不执行导出，通过时覆盖同一条回答。"""
    received = []

    async def create(self, **kwargs):
        """记录最终注入的 Word 正文，不能由规划模型替换。"""
        received.append(kwargs)
        return {'id': 'f' * 32, 'kind': 'word', 'title': kwargs['title'], 'filename': '笔记.docx',
                'media_type': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                'size_bytes': 100, 'download_url': '/api/v1/files/test/content'}

    async def search(self, session, user_id, question):
        """检查规划模型只提供检索问题，书籍证据来自用户的检索结果。"""
        assert question == 'Checkpoint 是什么？' and user_id == 7
        return [RetrievedChunk(BookChunk(id=SOURCE['chunk_id'], book_id=SOURCE['book_id'], user_id=7,
                                        title=SOURCE['title'], page=4, text=SOURCE['excerpt']), score=.9)]

    monkeypatch.setattr(KnowledgeService, 'search', search)
    monkeypatch.setattr(ArtifactService, 'create', create)
    model = ToolModel([call('generate_word', title='Checkpoint 笔记', content_source='knowledge',
                            question='Checkpoint 是什么？')], supported=supported)
    result = asyncio.run(build_chat_graph(model, InMemorySaver(), None).ainvoke(
        {'messages': [HumanMessage(content='请根据知识库生成 Checkpoint 学习 Word 文档')]},
        config={'configurable': {'thread_id': 'word-thread'}}, context=ChatContext(user_id=7)))
    assert len(result['messages']) == 2
    assert len(received) == int(supported)
    if supported:
        assert received[0]['content'] == 'Checkpoint 保存图状态 [1]。'
        assert received[0]['sources'][0]['chunk_id'] == SOURCE['chunk_id']
        assert result['messages'][-1].additional_kwargs['artifacts'][0]['kind'] == 'word'
    else:
        assert '没有足够依据' in result['messages'][-1].content
        assert not result['messages'][-1].additional_kwargs['artifacts']


def test_export_existing_answer_skips_retrieval_and_reuses_verified_text(monkeypatch):
    """导出旧答案不重新生成知识内容，也不重新请求检索模型服务。"""
    exported = AsyncMock(return_value={'id': 'a' * 32, 'kind': 'word', 'title': '笔记'})
    monkeypatch.setattr(ArtifactService, 'create', exported)
    monkeypatch.setattr(KnowledgeService, 'search', AsyncMock(side_effect=AssertionError('不应重新检索')))
    model = ToolModel([call('generate_word', title='笔记', content_source='latest_answer')])

    async def run():
        """先放入已核对历史，再连续两次导出，附件确认信息不会替换可信正文。"""
        graph = build_chat_graph(model, InMemorySaver(), None)
        config = {'configurable': {'thread_id': 'export'}}
        await graph.aupdate_state(config, {'messages': [HumanMessage(content='Checkpoint 是什么？'),
            AIMessage(content='保存图状态 [1]。', additional_kwargs={'sources': [SOURCE]})]}, as_node='save_memory')
        for _ in range(2):
            await graph.ainvoke({'messages': [HumanMessage(content='把刚才的回答导出为 Word')]},
                                config=config, context=ChatContext(user_id=7))
        assert exported.await_count == 2
        assert all(item.kwargs['content'] == '保存图状态 [1]。' for item in exported.await_args_list)
        assert model.counts == {'planner': 2}
    asyncio.run(run())


def test_word_without_verified_history_is_rejected_before_tool_execution(monkeypatch):
    """没有可导出的书籍回答时明确提示，不能生成空文档或从模型常识补正文。"""
    exported = AsyncMock()
    monkeypatch.setattr(ArtifactService, 'create', exported)
    model = ToolModel([call('generate_word', title='笔记', content_source='latest_answer')])
    with pytest.raises(ArtifactError, match='没有可导出的已核对回答'):
        asyncio.run(build_chat_graph(model, InMemorySaver(), None).ainvoke(
            {'messages': [HumanMessage(content='导出刚才的回答为 Word')]},
            config={'configurable': {'thread_id': 'missing-answer'}}, context=ChatContext(user_id=7)))
    assert exported.await_count == 0


def test_provider_failure_remains_upstream_error_without_retry(monkeypatch):
    """收费生成请求失败后不整节点重跑，并保留供应商错误类型供接口返回 502。"""
    exported = AsyncMock(side_effect=ArtifactProviderError('图片服务暂时不可用'))
    monkeypatch.setattr(ArtifactService, 'create', exported)
    model = ToolModel([call('generate_image', prompt='蓝色小猫插画', title='小猫')])
    with pytest.raises(ArtifactProviderError, match='图片服务暂时不可用'):
        asyncio.run(build_chat_graph(model, InMemorySaver(), None).ainvoke(
            {'messages': [HumanMessage(content='请生成一张小猫图片')]},
            config={'configurable': {'thread_id': 'failed-image'}}, context=ChatContext(user_id=7)))
    assert exported.await_count == 1


@pytest.mark.parametrize('invalid', ['wrong_host', 'redirect', 'too_large', 'invalid_image', 'provider_error', None])
def test_image_provider_contract_size_redirects_and_credentials(monkeypatch, invalid):
    """验证图片接口约定、体积上限、可信下载域名，以及下载不携带凭据。"""
    secret = 'dummy-image-key'
    monkeypatch.setattr(settings, 'IMAGE_API_KEY', secret)
    monkeypatch.setattr(settings, 'IMAGE_BASE_URL', 'https://api.example.com/v1')
    monkeypatch.setattr(settings, 'IMAGE_DOWNLOAD_HOSTS', 'images.example.com')
    monkeypatch.setattr('app.artifacts.images.socket.getaddrinfo', lambda *a, **kw: [(2, 1, 6, '', ('1.1.1.1', 443))])
    monkeypatch.setattr(settings, 'IMAGE_MAX_BYTES', 1024 if invalid == 'too_large' else 10000)
    calls = []

    def respond(request):
        """模拟生成接口及单独的图片下载请求，拒绝任何意外请求。"""
        calls.append(request)
        if request.method == 'POST':
            assert request.headers['authorization'] == 'Bearer ' + secret
            assert request.headers.get('x-enable-watermark') != '0'
            body = json.loads(request.content)
            assert body['image_size'] == '1024x1024' and 'batch_size' not in body
            if invalid == 'provider_error': return httpx.Response(401, json={'error': secret})
            host = '127.0.0.1' if invalid == 'wrong_host' else 'images.example.com'
            return httpx.Response(200, json={'images': [{'url': f'https://{host}/output.png'}]})
        assert 'authorization' not in request.headers
        if invalid == 'redirect': return httpx.Response(302, headers={'location': 'http://127.0.0.1/private'})
        data = b'x' * 2048 if invalid == 'too_large' else b'not-an-image' if invalid == 'invalid_image' else png_bytes()
        return httpx.Response(200, content=data)

    original = httpx.AsyncClient
    monkeypatch.setattr('app.artifacts.images.httpx.AsyncClient',
                        lambda **kw: original(transport=httpx.MockTransport(respond), **kw))
    if invalid:
        with pytest.raises(ArtifactProviderError) as raised:
            asyncio.run(generate_image_bytes('蓝色小猫', 'square'))
        assert secret not in str(raised.value)
    else:
        data, extension, mime = asyncio.run(generate_image_bytes('蓝色小猫', 'square'))
        assert data == png_bytes() and extension == 'png' and mime == 'image/png'
    assert len(calls) <= 2


def test_allowed_image_host_resolving_to_private_ip_is_rejected(monkeypatch):
    """即使域名匹配配置，也不能连接到内网地址。"""
    monkeypatch.setattr(settings, 'IMAGE_DOWNLOAD_HOSTS', 'images.example.com')
    monkeypatch.setattr('app.artifacts.images.socket.getaddrinfo', lambda *a, **kw: [(2, 1, 6, '', ('127.0.0.1', 443))])
    with pytest.raises(ArtifactProviderError):
        asyncio.run(validate_image_url('https://images.example.com/a.png'))


def test_generated_file_is_private_and_tool_result_is_idempotent(tmp_path, monkeypatch):
    """用真实关系库和 HTTP 接口验证文件下载、用户隔离、路径边界与同请求复用。"""
    async def run():
        """所有生成文件和数据库都写临时目录，不改变生产会话。"""
        engine = create_async_engine(f'sqlite+aiosqlite:///{(tmp_path / "artifacts.db").as_posix()}')
        monkeypatch.setattr('app.services.artifact_service.engine', engine)
        monkeypatch.setattr(settings, 'ARTIFACT_STORAGE_DIR', tmp_path / 'generated')
        async with engine.begin() as connection:
            await connection.run_sync(SQLModel.metadata.create_all)

        async def session_override():
            """为认证、下载接口提供同一测试数据库的独立请求会话。"""
            async with AsyncSession(engine) as session:
                yield session

        app.dependency_overrides[get_session] = session_override
        try:
            async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
                users = []
                for name in ('alice', 'bob'):
                    response = await client.post('/api/v1/auth/register', json={
                        'username': name, 'email': f'{name}@example.com', 'password': 'pass12345'})
                    assert response.status_code == 201, response.text
                    users.append({'Authorization': 'Bearer ' + response.json()['access_token']})
                args = {'title': '学习笔记', 'user_id': 1, 'thread_id': 'owned-thread',
                        'request_key': 'same-request', 'content': '保存图状态 [1]。', 'sources': [SOURCE]}
                first = await generate_word.ainvoke(args)
                second = await generate_word.ainvoke(args)
                assert first == second and 'storage_path' not in first
                content = await client.get(first['download_url'], headers=users[0])
                assert content.status_code == 200
                assert Document(BytesIO(content.content)).paragraphs[0].text == '学习笔记'
                assert (await client.get(first['download_url'], headers=users[1])).status_code == 404
                assert (await client.get(first['download_url'])).status_code in (401, 403)
                async with AsyncSession(engine) as session:
                    rows = list((await session.exec(select(GeneratedArtifact))).all())
                    assert len(rows) == 1
                    assert ArtifactService.storage_path(rows[0]).is_file()
                    rows[0].storage_path = str(tmp_path / 'outside.txt')
                    session.add(rows[0])
                    await session.commit()
                (tmp_path / 'outside.txt').write_text('不得读取', encoding='utf-8')
                assert (await client.get(first['download_url'], headers=users[0])).status_code == 404
        finally:
            app.dependency_overrides.clear()
            await engine.dispose()
    asyncio.run(run())


def test_word_tool_sse_history_and_authenticated_download(tmp_path, monkeypatch):
    """完整跑通登录、会话、模型选工具、真实 Word、SSE、历史恢复和认证下载。"""
    async def run():
        """使用隔离关系库与真实图，保留实际接口及文件生成代码。"""
        engine = create_async_engine(f'sqlite+aiosqlite:///{(tmp_path / "chat.db").as_posix()}')
        monkeypatch.setattr('app.services.artifact_service.engine', engine)
        monkeypatch.setattr(settings, 'ARTIFACT_STORAGE_DIR', tmp_path / 'generated')
        async with engine.begin() as connection:
            await connection.run_sync(SQLModel.metadata.create_all)

        async def session_override():
            """让认证、会话和文件接口连接隔离数据库。"""
            async with AsyncSession(engine) as session:
                yield session

        app.dependency_overrides[get_session] = session_override
        original_graph = getattr(app.state, 'chat_graph', None)
        try:
            async with AsyncClient(transport=ASGITransport(app=app), base_url='http://test') as client:
                registration = await client.post('/api/v1/auth/register', json={
                    'username': 'reader', 'email': 'reader@example.com', 'password': 'pass12345'})
                headers = {'Authorization': 'Bearer ' + registration.json()['access_token']}
                created = await client.post('/api/v1/conversations', headers=headers, json={'title': '导出测试'})
                conversation_id = created.json()['id']
                async with AsyncSession(engine) as session:
                    conversation = await session.get(Conversation, conversation_id)
                    thread_id = conversation.thread_id
                model = ToolModel([call('generate_word', title='学习笔记', content_source='latest_answer')])
                graph = build_chat_graph(model, InMemorySaver(), None)
                app.state.chat_graph = graph
                await graph.aupdate_state({'configurable': {'thread_id': thread_id}}, {'messages': [
                    HumanMessage(content='Checkpoint 是什么？'),
                    AIMessage(content='保存图状态 [1]。', additional_kwargs={'sources': [SOURCE]})]}, as_node='save_memory')
                response = await client.post('/api/v1/chat/stream', headers=headers,
                    json={'conversation_id': conversation_id, 'message': '把刚才的回答导出为 Word'})
                assert response.status_code == 200
                events = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith('data: ')]
                assert events[-1]['type'] == 'done'
                artifacts = next(event['artifacts'] for event in events if event['type'] == 'artifacts')
                assert artifacts[0]['kind'] == 'word'
                assert sum(event['type'] == 'answer' for event in events) == 1
                history = await client.get(f'/api/v1/chat/{conversation_id}/history', headers=headers)
                assert len(history.json()['messages']) == 4
                assert history.json()['messages'][-1]['artifacts'] == artifacts
                download = await client.get(artifacts[0]['download_url'], headers=headers)
                assert download.status_code == 200
                document = Document(BytesIO(download.content))
                assert any('保存图状态' in paragraph.text for paragraph in document.paragraphs)
        finally:
            app.dependency_overrides.clear()
            if original_graph is None:
                del app.state.chat_graph
            else:
                app.state.chat_graph = original_graph
            await engine.dispose()
    asyncio.run(run())
