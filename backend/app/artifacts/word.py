"""把已有正文转换成 Word，保留 Markdown 标题、列表、代码和表格结构。"""

import io
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from markdown_it import MarkdownIt

from app.artifacts.errors import ArtifactError


def clean_xml_text(text: str) -> str:
    """移除 XML 禁止的控制字符，避免 PDF 摘录导致 Word 文件写入失败。"""
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]", "", text)


def add_inline(paragraph, token) -> None:
    """把行内加粗、斜体、代码和换行转换为 Word 文本格式。"""
    bold = italic = False
    for child in token.children or []:
        if child.type == "strong_open": bold = True
        elif child.type == "strong_close": bold = False
        elif child.type == "em_open": italic = True
        elif child.type == "em_close": italic = False
        elif child.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif child.type in ("text", "code_inline"):
            run = paragraph.add_run(clean_xml_text(child.content))
            run.bold, run.italic = bold, italic
            if child.type == "code_inline": run.font.name = "Consolas"


def word_bytes(title: str, content: str, sources: list[dict]) -> bytes:
    """生成中文 A4 学习文档；仅排版后端提供的正文，不另行调用 LLM 编造内容。"""
    if not content.strip() or len(content) > 60000:
        raise ArtifactError("Word 正文为空或过长，请缩小导出范围。")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal = document.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    document.add_paragraph(clean_xml_text(title), style="Title")
    tokens = MarkdownIt("commonmark").enable("table").parse(content)
    lists, paragraph = [], None
    table = row = cell = None
    for token in tokens:
        kind = token.type
        if kind in ("bullet_list_open", "ordered_list_open"):
            lists.append("List Bullet" if kind == "bullet_list_open" else "List Number")
        elif kind in ("bullet_list_close", "ordered_list_close"):
            lists.pop()
        elif kind == "heading_open":
            paragraph = document.add_paragraph(style=f"Heading {min(int(token.tag[1:]), 3)}")
        elif kind == "paragraph_open":
            paragraph = document.add_paragraph(style=lists[-1] if lists else "Normal")
        elif kind == "inline":
            if paragraph is None: paragraph = document.add_paragraph()
            add_inline(paragraph, token)
        elif kind in ("heading_close", "paragraph_close"):
            paragraph = None
        elif kind in ("fence", "code_block"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(clean_xml_text(token.content.rstrip()))
            run.font.name, run.font.size = "Consolas", Pt(9)
            paragraph = None
        elif kind == "table_open":
            table = document.add_table(rows=0, cols=0)
            table.style = "Table Grid"
        elif kind == "tr_open":
            row = table.add_row()
            cell = 0
        elif kind in ("th_open", "td_open"):
            if cell >= len(table.columns): table.add_column(Cm(15 / max(cell + 1, 1)))
            target = row.cells[cell]
            paragraph = target.paragraphs[0]
            if kind == "th_open":
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EDEDED")
                target._tc.get_or_add_tcPr().append(shade)
            cell += 1
        elif kind in ("th_close", "td_close"):
            paragraph = None
        elif kind == "table_close":
            width = Cm(16.6 / max(len(table.columns), 1))
            for column in table.columns: column.width = width
            document.add_paragraph()
    if sources:
        document.add_heading("参考文献", level=1)
        for index, source in enumerate(sources, 1):
            page = source.get("page", 0)
            location = f"第 {page} 页" if page > 0 else f"第 {-page} 段"
            document.add_paragraph(clean_xml_text(f"[{index}] {source.get('title', '')} · {location}"))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()
